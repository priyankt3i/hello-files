"""Exact spreadsheet calculations.

Spreadsheets are re-read from disk (not from the search index, which only holds row text),
loaded as typed tables into an in-memory SQLite database, and queried with a single,
locked-down SELECT. Nothing here touches the index or the user's files beyond reading them.

Two entry points, both exposed over the worker protocol:

* ``spreadsheet_schema``: tables, columns, inferred types, row counts and sample rows.
* ``spreadsheet_query``: run one read-only SELECT over those tables.

Safety model for the query (the SQL is written by a language model and is untrusted):
a SQLite authorizer allows only SELECT, reads of the loaded tables and a whitelist of pure
functions; the statement must be a single SELECT/WITH; a progress handler aborts runaway
queries; results are row-capped; the database is in memory and PRAGMA query_only is set.
"""

from __future__ import annotations

import csv
import io
import math
import re
import sqlite3
import time
import unicodedata
import warnings
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path
from typing import Any

from extractors import (
    HEADER_SCAN_LIMIT,
    detect_header_row,
    load_xlrd,
    normalize_spreadsheet_cell,
    sanitize_header_values,
)
from index_store import read_documents

SPREADSHEET_PARSERS = {"spreadsheet", "spreadsheet-xls", "csv"}
SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv"}

DEFAULT_MAX_ROWS = 200
HARD_MAX_ROWS = 1000
QUERY_TIMEOUT_SECONDS = 5.0
MAX_SQL_LENGTH = 8000
MAX_SHEET_ROWS = 500_000
MAX_SHEET_COLUMNS = 256
MAX_SCHEMA_DOCUMENTS = 25
SAMPLE_ROW_COUNT = 3
CATEGORY_MAX_DISTINCT = 12
CATEGORY_MAX_LENGTH = 40
CACHE_CELL_BUDGET = 30_000_000
MAX_SAFE_INTEGER = 2**53

ALLOWED_FUNCTIONS = frozenset(
    {
        # aggregates
        "count", "sum", "avg", "min", "max", "total", "group_concat",
        # scalar
        "round", "abs", "coalesce", "nullif", "ifnull", "iif", "lower", "upper", "trim", "ltrim", "rtrim",
        "substr", "substring", "length", "replace", "instr", "typeof", "like", "glob", "printf", "format",
        "floor", "ceil", "ceiling", "sqrt", "power", "sign",
        # dates
        "strftime", "date", "datetime", "time", "julianday",
        # window functions
        "row_number", "rank", "dense_rank", "lag", "lead", "first_value", "last_value", "ntile",
    }
)

# Names that would need quoting; a trailing underscore keeps generated SQL simple.
SQL_KEYWORDS = frozenset(
    {
        "all", "and", "as", "asc", "between", "by", "case", "check", "collate", "column", "create", "cross",
        "default", "delete", "desc", "distinct", "drop", "else", "end", "except", "exists", "foreign", "from",
        "full", "glob", "group", "having", "in", "index", "inner", "insert", "intersect", "into", "is", "join",
        "key", "left", "like", "limit", "match", "natural", "not", "null", "offset", "on", "or", "order", "outer",
        "primary", "references", "right", "select", "set", "table", "then", "to", "union", "unique", "update",
        "using", "values", "view", "when", "where", "with", "window", "over", "partition", "range", "rows",
    }
)

NULL_TOKENS = frozenset({"n/a", "na", "null", "none", "nan", "-", "--", "#n/a", "#null!", "#div/0!", "#value!", "#ref!", "#name?", "#num!"})

KIND_NULL = "null"
KIND_TOKEN = "token"
KIND_INT = "int"
KIND_REAL = "real"
KIND_DATE = "date"
KIND_TEXT = "text"


class SpreadsheetQueryError(Exception):
    """A problem with the request or the SQL that the caller can show to the model."""


# ---------------------------------------------------------------------------
# Value classification and coercion
# ---------------------------------------------------------------------------

_ISO_DATE_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2})(?:\.\d+)?)?)?(?:Z)?$")
_SLASH_DATE_RE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?(?:\s*([AaPp][Mm]))?)?$")
_YMD_SLASH_RE = re.compile(r"^(\d{4})/(\d{1,2})/(\d{1,2})$")
_NUMBER_RE = re.compile(r"^(\()?([+-])?\s*[$€£]?\s*(\d{1,3}(?:,\d{3})+|\d*)(?:\.(\d+))?(?:[eE]([+-]?\d+))?(\))?$")


def _format_datetime(value: datetime) -> str:
    if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
        return value.strftime("%Y-%m-%d")
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _make_datetime(year: int, month: int, day: int, hour: int = 0, minute: int = 0, second: int = 0) -> str | None:
    try:
        return _format_datetime(datetime(year, month, day, hour, minute, second))
    except ValueError:
        return None


def _classify_string(text: str) -> tuple[str, Any]:
    lowered = text.lower()
    if lowered in NULL_TOKENS:
        return KIND_TOKEN, None
    first = text[0]
    if not (first.isdigit() or first in "+-.($€£"):
        return KIND_TEXT, text

    match = _NUMBER_RE.match(text)
    if match:
        open_paren, sign, integer_part, decimals, exponent, close_paren = match.groups()
        digits = integer_part.replace(",", "")
        if bool(open_paren) == bool(close_paren) and (digits or decimals):
            # "007" and long digit strings are identifiers (zip codes, policy numbers), not numbers.
            if (len(digits) > 1 and digits.startswith("0")) or len(digits) > 15 or (decimals and len(decimals) > 15):
                return KIND_TEXT, text
            number_text = f"{digits or '0'}{'.' + decimals if decimals else ''}{'e' + exponent if exponent else ''}"
            negative = sign == "-" or bool(open_paren)
            if decimals or exponent:
                value = float(number_text)
                if not math.isfinite(value):
                    return KIND_TEXT, text
                return KIND_REAL, -value if negative else value
            integer = int(number_text)
            return KIND_INT, -integer if negative else integer

    iso = _ISO_DATE_RE.match(text)
    if iso:
        year, month, day, hour, minute, second = iso.groups()
        formatted = _make_datetime(int(year), int(month), int(day), int(hour or 0), int(minute or 0), int(second or 0))
        if formatted:
            return KIND_DATE, formatted
        return KIND_TEXT, text

    slash = _SLASH_DATE_RE.match(text)
    if slash:
        first_part, second_part, year_text, hour, minute, second, meridiem = slash.groups()
        month, day = int(first_part), int(second_part)
        if month > 12 and day <= 12:  # clearly day-first
            month, day = day, month
        year = int(year_text)
        if len(year_text) == 2:
            year += 2000 if year < 70 else 1900
        hour_value = int(hour or 0)
        if meridiem:
            hour_value = hour_value % 12 + (12 if meridiem.lower() == "pm" else 0)
        formatted = _make_datetime(year, month, day, hour_value, int(minute or 0), int(second or 0))
        if formatted:
            return KIND_DATE, formatted
        return KIND_TEXT, text

    year_first = _YMD_SLASH_RE.match(text)
    if year_first:
        formatted = _make_datetime(int(year_first.group(1)), int(year_first.group(2)), int(year_first.group(3)))
        if formatted:
            return KIND_DATE, formatted
    return KIND_TEXT, text


def _classify(value: Any) -> tuple[str, Any]:
    """Return (kind, coerced value) for one raw cell. Text kinds keep the original text."""
    if value is None:
        return KIND_NULL, None
    cell_type = type(value)
    if cell_type is str:
        stripped = value.strip()
        if not stripped:
            return KIND_NULL, None
        return _classify_string(stripped)
    if cell_type is bool:
        return KIND_INT, int(value)
    if cell_type is int:
        if abs(value) >= MAX_SAFE_INTEGER:
            return KIND_TEXT, str(value)
        return KIND_INT, value
    if cell_type is float:
        if not math.isfinite(value):
            return KIND_TOKEN, None
        if value.is_integer() and abs(value) < MAX_SAFE_INTEGER:
            return KIND_INT, int(value)
        return KIND_REAL, value
    if isinstance(value, datetime):
        return KIND_DATE, _format_datetime(value)
    if isinstance(value, date):
        return KIND_DATE, value.strftime("%Y-%m-%d")
    if isinstance(value, dt_time):
        return KIND_TEXT, value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        return KIND_TEXT, str(value)
    return KIND_TEXT, str(value)


def _original_text(value: Any) -> str | None:
    """The cell as the user would see it, for text columns and the *_raw companions."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return str(int(value)) if value.is_integer() and abs(value) < MAX_SAFE_INTEGER else format(value, ".15g")
    if isinstance(value, datetime):
        return _format_datetime(value)
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt_time):
        return value.strftime("%H:%M:%S")
    return str(value)


# ---------------------------------------------------------------------------
# Loaded sheets
# ---------------------------------------------------------------------------


@dataclass
class Column:
    name: str
    original: str
    type: str  # integer | real | text | date
    values: list[Any]
    non_null: int
    raw_name: str | None = None
    raw_values: list[Any] | None = None
    coerce_failures: int = 0
    categories: list[str] | None = None


@dataclass
class LoadedSheet:
    sheet: str
    header_row: int | None
    row_numbers: list[int]
    columns: list[Column]

    @property
    def row_count(self) -> int:
        return len(self.row_numbers)


@dataclass
class LoadedWorkbook:
    sheets: list[LoadedSheet]
    cells: int = 0


@dataclass
class SheetTable:
    document: str
    sheet: str
    table: str
    loaded: LoadedSheet


_CACHE: "OrderedDict[str, tuple[tuple[int, int], LoadedWorkbook]]" = OrderedDict()


def clear_spreadsheet_cache() -> None:
    _CACHE.clear()


def _sanitize_identifier(text: str, fallback: str, limit: int = 48) -> str:
    replaced = text.replace("%", " pct ").replace("#", " num ").replace("&", " and ").replace("$", " usd ")
    ascii_text = unicodedata.normalize("NFKD", replaced).encode("ascii", "ignore").decode("ascii").lower()
    name = re.sub(r"[^0-9a-z]+", "_", ascii_text).strip("_")[:limit].strip("_")
    if not name:
        name = fallback
    if name[0].isdigit():
        name = f"c_{name}"
    return name


def _unique_name(name: str, used: set[str]) -> str:
    candidate = name
    counter = 2
    while candidate.lower() in used:
        candidate = f"{name}_{counter}"
        counter += 1
    used.add(candidate.lower())
    return candidate


def _trim_raw(values: list[Any]) -> list[Any]:
    end = len(values)
    while end > 0:
        value = values[end - 1]
        if value is None or (isinstance(value, str) and not value.strip()):
            end -= 1
        else:
            break
    return values[:end]


def _read_xlsx_rows(path: Path) -> list[tuple[str, list[tuple[int, list[Any]]]]]:
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Conditional Formatting extension is not supported and will be removed",
            category=UserWarning,
            module=r"openpyxl\.worksheet\._reader",
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[tuple[str, list[tuple[int, list[Any]]]]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[int, list[Any]]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = _trim_raw(list(row))
                if values:
                    rows.append((row_index, values))
                    if len(rows) > MAX_SHEET_ROWS:
                        raise SpreadsheetQueryError(
                            f"Sheet '{sheet.title}' has more than {MAX_SHEET_ROWS:,} rows, which is too large for exact calculation."
                        )
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    return sheets


def _read_xls_rows(path: Path) -> list[tuple[str, list[tuple[int, list[Any]]]]]:
    xlrd = load_xlrd()
    if xlrd is None:
        raise SpreadsheetQueryError("Legacy .xls support requires xlrd.")
    from xlrd.xldate import xldate_as_datetime

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    sheets: list[tuple[str, list[tuple[int, list[Any]]]]] = []
    try:
        for sheet in workbook.sheets():
            if sheet.nrows > MAX_SHEET_ROWS:
                raise SpreadsheetQueryError(
                    f"Sheet '{sheet.name}' has more than {MAX_SHEET_ROWS:,} rows, which is too large for exact calculation."
                )
            rows: list[tuple[int, list[Any]]] = []
            for row_index in range(sheet.nrows):
                values: list[Any] = []
                for cell in sheet.row(row_index):
                    ctype, value = cell.ctype, cell.value
                    if ctype in (0, 6):
                        values.append(None)
                    elif ctype == 3:
                        try:
                            converted = xldate_as_datetime(value, workbook.datemode)
                            values.append(converted.time() if 0 <= value < 1 else converted)
                        except (ValueError, OverflowError):
                            values.append(value)
                    elif ctype == 4:
                        values.append(bool(value))
                    elif ctype == 5:
                        values.append("#N/A")
                    else:
                        values.append(value)
                trimmed = _trim_raw(values)
                if trimmed:
                    rows.append((row_index + 1, trimmed))
            sheets.append((sheet.name, rows))
            workbook.unload_sheet(sheet.name)
    finally:
        workbook.release_resources()
    return sheets


def _read_csv_rows(path: Path) -> list[tuple[str, list[tuple[int, list[Any]]]]]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    first_line = text.split("\n", 1)[0]
    delimiter = max(",;\t|", key=lambda candidate: (first_line.count(candidate), candidate == ","))
    if first_line.count(delimiter) == 0:
        delimiter = ","
    csv.field_size_limit(10_000_000)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[tuple[int, list[Any]]] = []
    try:
        for record in reader:
            values = _trim_raw(list(record))
            if values:
                rows.append((reader.line_num, values))
                if len(rows) > MAX_SHEET_ROWS:
                    raise SpreadsheetQueryError(f"{path.name} has more than {MAX_SHEET_ROWS:,} rows, which is too large for exact calculation.")
    except csv.Error as exc:
        raise SpreadsheetQueryError(f"{path.name} could not be read as CSV: {exc}") from exc
    return [(path.name, rows)]


def _build_sheet(title: str, rows: list[tuple[int, list[Any]]]) -> LoadedSheet | None:
    if not rows:
        return None
    probe = [(number, [normalize_spreadsheet_cell(value) for value in values]) for number, values in rows[:HEADER_SCAN_LIMIT]]
    header_index = detect_header_row(probe)
    if header_index is not None:
        header_row_number = rows[header_index][0]
        headers = sanitize_header_values(probe[header_index][1])
        data_rows = rows[header_index + 1 :]
    else:
        header_row_number = None
        headers = []
        data_rows = rows

    width = max([len(headers)] + [len(values) for _, values in data_rows])
    if width > MAX_SHEET_COLUMNS:
        raise SpreadsheetQueryError(f"Sheet '{title}' has more than {MAX_SHEET_COLUMNS} columns, which is too large for exact calculation.")
    originals = [headers[index] if index < len(headers) else f"Column {index + 1}" for index in range(width)]

    used: set[str] = {"_row"}
    names: list[str] = []
    for index, original in enumerate(originals, start=1):
        name = _sanitize_identifier(original, f"column_{index}")
        if name in SQL_KEYWORDS:
            name = f"{name}_"
        names.append(_unique_name(name, used))

    row_numbers = [number for number, _ in data_rows]
    columns: list[Column] = []
    for index in range(width):
        raw_values = [values[index] if index < len(values) else None for _, values in data_rows]
        columns.append(_build_column(names[index], originals[index], raw_values, used))
    return LoadedSheet(sheet=title, header_row=header_row_number, row_numbers=row_numbers, columns=columns)


def _build_column(name: str, original: str, raw_values: list[Any], used: set[str]) -> Column:
    classified = [_classify(value) for value in raw_values]
    counts: dict[str, int] = {}
    for kind, _ in classified:
        counts[kind] = counts.get(kind, 0) + 1
    non_blank = len(classified) - counts.get(KIND_NULL, 0)
    considered = non_blank - counts.get(KIND_TOKEN, 0)
    numeric = counts.get(KIND_INT, 0) + counts.get(KIND_REAL, 0)
    dates = counts.get(KIND_DATE, 0)

    if considered > 0 and numeric / considered >= 0.9 and numeric >= dates:
        column_type = "real" if counts.get(KIND_REAL, 0) > 0 else "integer"
        accepted = {KIND_INT, KIND_REAL}
    elif considered > 0 and dates / considered >= 0.9:
        column_type = "date"
        accepted = {KIND_DATE}
    else:
        column_type = "text"
        accepted = set()

    if column_type == "text":
        values = [_original_text(raw) for raw in raw_values]
        non_null = sum(1 for value in values if value is not None)
        return Column(name=name, original=original, type="text", values=values, non_null=non_null, categories=_categories(values, non_null))

    values: list[Any] = []
    raw_text: list[str | None] = []
    failures = 0
    for raw, (kind, coerced) in zip(raw_values, classified):
        if kind in accepted:
            values.append(coerced)
            raw_text.append(None)
        else:
            values.append(None)
            if kind in (KIND_NULL, KIND_TOKEN):
                raw_text.append(_original_text(raw) if kind == KIND_TOKEN else None)
            else:
                failures += 1
                raw_text.append(_original_text(raw))
    column = Column(name=name, original=original, type=column_type, values=values, non_null=sum(1 for value in values if value is not None))
    if any(text is not None for text in raw_text):
        column.raw_name = _unique_name(f"{name}_raw", used)
        column.raw_values = raw_text
        column.coerce_failures = failures
    return column


def _categories(values: list[Any], non_null: int) -> list[str] | None:
    if non_null < 2:
        return None
    distinct = set(value for value in values if value is not None)
    if len(distinct) > min(CATEGORY_MAX_DISTINCT, non_null // 2):
        return None
    if any(len(value) > CATEGORY_MAX_LENGTH for value in distinct):
        return None
    return sorted(distinct, key=lambda item: (item.lower(), item))


def _load_workbook(path: Path) -> LoadedWorkbook:
    stat = path.stat()
    signature = (stat.st_mtime_ns, stat.st_size)
    key = str(path)
    cached = _CACHE.get(key)
    if cached and cached[0] == signature:
        _CACHE.move_to_end(key)
        return cached[1]
    _CACHE.pop(key, None)

    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        raw_sheets = _read_xlsx_rows(path)
    elif suffix == ".xls":
        raw_sheets = _read_xls_rows(path)
    elif suffix == ".csv":
        raw_sheets = _read_csv_rows(path)
    else:
        raise SpreadsheetQueryError(f"{path.name} is not a spreadsheet.")

    sheets = [sheet for sheet in (_build_sheet(title, rows) for title, rows in raw_sheets) if sheet is not None]
    workbook = LoadedWorkbook(sheets=sheets, cells=sum(sheet.row_count * max(len(sheet.columns), 1) for sheet in sheets))
    _CACHE[key] = (signature, workbook)
    total = sum(entry[1].cells for entry in _CACHE.values())
    while total > CACHE_CELL_BUDGET and len(_CACHE) > 1:
        _, (_, evicted) = _CACHE.popitem(last=False)
        total -= evicted.cells
    return workbook


# ---------------------------------------------------------------------------
# Resolving documents and naming tables
# ---------------------------------------------------------------------------


def _norm_key(value: str) -> str:
    return value.replace("\\", "/")


def _resolve_documents(root_path: Path, index_path: Path, document_ids: list[str] | None) -> tuple[list[dict[str, Any]], bool]:
    documents = [
        document
        for document in read_documents(index_path)
        if document.get("parser") in SPREADSHEET_PARSERS or Path(document["relativePath"]).suffix.lower() in SPREADSHEET_SUFFIXES
    ]
    capped = False
    if document_ids is None:
        if len(documents) > MAX_SCHEMA_DOCUMENTS:
            documents = documents[:MAX_SCHEMA_DOCUMENTS]
            capped = True
        return documents, capped
    by_id = {_norm_key(document["documentId"]): document for document in documents}
    selected: list[dict[str, Any]] = []
    seen: set[str] = set()
    for document_id in document_ids:
        key = _norm_key(str(document_id))
        if key in seen:
            continue
        seen.add(key)
        if key not in by_id:
            raise SpreadsheetQueryError(f"'{document_id}' is not an indexed spreadsheet.")
        selected.append(by_id[key])
    if len(selected) > MAX_SCHEMA_DOCUMENTS:
        raise SpreadsheetQueryError(f"At most {MAX_SCHEMA_DOCUMENTS} spreadsheets can be used in one calculation.")
    return selected, False


def _source_path(root_path: Path, relative_path: str) -> Path:
    root = root_path.resolve()
    candidate = (root / relative_path).resolve()
    if candidate != root and root not in candidate.parents:
        raise SpreadsheetQueryError(f"'{relative_path}' is outside the indexed folder.")
    if not candidate.is_file():
        raise SpreadsheetQueryError("File not found on disk (it may have been moved or deleted since the last sync).")
    return candidate


def _table_base(relative_path: str, sheet: str, is_csv: bool) -> str:
    path = Path(relative_path)
    stem = _sanitize_identifier(path.stem, "sheet", 28)
    extension = path.suffix.lower().lstrip(".") or "file"
    if is_csv:
        return f"{stem}_{extension}"
    return f"{stem}_{extension}__{_sanitize_identifier(sheet, 'sheet', 28)}"


def _load_tables(root_path: Path, index_path: Path, document_ids: list[str] | None) -> tuple[list[SheetTable], list[dict[str, str]], bool]:
    """Load every requested spreadsheet. Files that cannot be read are reported, not fatal."""
    documents, capped = _resolve_documents(root_path, index_path, document_ids)
    documents = sorted(documents, key=lambda item: _norm_key(item["relativePath"]))
    tables: list[SheetTable] = []
    skipped: list[dict[str, str]] = []
    used: set[str] = set()
    for document in documents:
        relative_path = document["relativePath"]
        try:
            path = _source_path(root_path, relative_path)
            workbook = _load_workbook(path)
        except SpreadsheetQueryError as exc:
            skipped.append({"document": relative_path, "reason": str(exc)})
            continue
        except Exception as exc:  # corrupt or password-protected files
            skipped.append({"document": relative_path, "reason": f"Could not read the file: {str(exc).strip() or exc.__class__.__name__}"})
            continue
        is_csv = path.suffix.lower() == ".csv"
        for sheet in workbook.sheets:
            table = _unique_name(_table_base(relative_path, sheet.sheet, is_csv), used)
            tables.append(SheetTable(document=relative_path, sheet=sheet.sheet, table=table, loaded=sheet))
    return tables, skipped, capped


# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool)):
        return value
    if isinstance(value, int):
        return value if abs(value) < MAX_SAFE_INTEGER else str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return round(value, 6)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return f"<{len(bytes(value))} bytes>"
    return str(value)


def spreadsheet_schema(root_path: Path, index_path: Path, document_ids: list[str] | None = None) -> dict[str, Any]:
    tables, skipped, capped = _load_tables(root_path, index_path, document_ids)
    documents: list[dict[str, Any]] = []
    by_document: dict[str, dict[str, Any]] = {}
    for item in tables:
        entry = by_document.get(item.document)
        if entry is None:
            entry = {"document": item.document, "sheets": []}
            by_document[item.document] = entry
            documents.append(entry)
        sheet = item.loaded
        columns = []
        for column in sheet.columns:
            description: dict[str, Any] = {
                "name": column.name,
                "original": column.original,
                "type": column.type,
                "nonNull": column.non_null,
            }
            if column.raw_name:
                description["rawColumn"] = column.raw_name
                description["coerceFailures"] = column.coerce_failures
            if column.categories:
                description["values"] = column.categories
            columns.append(description)
        sample_count = min(SAMPLE_ROW_COUNT, sheet.row_count)
        entry["sheets"].append(
            {
                "sheet": item.sheet,
                "table": item.table,
                "rowCount": sheet.row_count,
                "headerRow": sheet.header_row,
                "columns": columns,
                "sampleRows": [[_json_value(column.values[row]) for column in sheet.columns] for row in range(sample_count)],
            }
        )
    return {"documents": documents, "skipped": skipped, "truncatedDocuments": capped}


# ---------------------------------------------------------------------------
# Query
# ---------------------------------------------------------------------------


def _build_database(tables: list[SheetTable]) -> sqlite3.Connection:
    connection = sqlite3.connect(":memory:", check_same_thread=False)
    for item in tables:
        sheet = item.loaded
        definitions = ['"_row" INTEGER']
        series: list[list[Any]] = [sheet.row_numbers]
        for column in sheet.columns:
            sql_type = {"integer": "INTEGER", "real": "REAL"}.get(column.type, "TEXT")
            definitions.append(f'"{column.name}" {sql_type}')
            series.append(column.values)
            if column.raw_name and column.raw_values is not None:
                definitions.append(f'"{column.raw_name}" TEXT')
                series.append(column.raw_values)
        connection.execute(f'CREATE TABLE "{item.table}" ({", ".join(definitions)})')
        placeholders = ", ".join("?" for _ in series)
        connection.executemany(f'INSERT INTO "{item.table}" VALUES ({placeholders})', zip(*series))
    connection.commit()
    return connection


def _describe_tables(tables: list[SheetTable], limit: int = 1800) -> str:
    parts = []
    for item in tables:
        names = ", ".join(column.name for column in item.loaded.columns)
        parts.append(f"{item.table}({names})")
    text = "; ".join(parts)
    return text if len(text) <= limit else text[: limit - 3] + "..."


def _lock_down(connection: sqlite3.Connection, allowed_tables: dict[str, SheetTable], denied: list[str], used: dict[str, SheetTable], deadline: list[float]) -> None:
    connection.execute("PRAGMA query_only = ON")
    for constant, value in (("SQLITE_LIMIT_LENGTH", 2_000_000), ("SQLITE_LIMIT_SQL_LENGTH", MAX_SQL_LENGTH * 2), ("SQLITE_LIMIT_EXPR_DEPTH", 200)):
        limit_id = getattr(sqlite3, constant, None)
        if limit_id is not None and hasattr(connection, "setlimit"):
            connection.setlimit(limit_id, value)
    # Double-quoted names that match no column must be an error, not a silent string literal.
    if hasattr(connection, "setconfig"):
        for constant in ("SQLITE_DBCONFIG_DQS_DML", "SQLITE_DBCONFIG_DQS_DDL"):
            config_id = getattr(sqlite3, constant, None)
            if config_id is not None:
                connection.setconfig(config_id, False)

    def authorizer(action: int, arg1: str | None, arg2: str | None, db_name: str | None, source: str | None) -> int:
        if action == sqlite3.SQLITE_SELECT:
            return sqlite3.SQLITE_OK
        if action == sqlite3.SQLITE_READ:
            table = allowed_tables.get((arg1 or "").lower())
            if table is None:
                denied.append(f"Access to '{arg1}' is not allowed. Only the spreadsheet tables can be read (recursive queries are not supported).")
                return sqlite3.SQLITE_DENY
            used[table.table] = table
            return sqlite3.SQLITE_OK
        if action == sqlite3.SQLITE_FUNCTION:
            name = (arg2 or "").lower()
            if name in ALLOWED_FUNCTIONS:
                return sqlite3.SQLITE_OK
            denied.append(f"The function '{arg2}' is not allowed.")
            return sqlite3.SQLITE_DENY
        denied.append("Only read-only SELECT queries are allowed.")
        return sqlite3.SQLITE_DENY

    def progress() -> int:
        return 1 if time.monotonic() > deadline[0] else 0

    connection.set_authorizer(authorizer)
    connection.set_progress_handler(progress, 2000)


def _prepare_sql(sql: Any) -> str:
    if not isinstance(sql, str) or not sql.strip():
        raise SpreadsheetQueryError("The SQL query is empty.")
    text = sql.strip()
    if len(text) > MAX_SQL_LENGTH:
        raise SpreadsheetQueryError(f"The SQL query is too long (limit {MAX_SQL_LENGTH} characters).")
    text = text.rstrip(";").rstrip()
    stripped = re.sub(r"^(?:\s+|--[^\n]*(?:\n|$)|/\*.*?\*/)+", "", text, flags=re.S)
    if not re.match(r"(?i)(select|with)\b", stripped):
        raise SpreadsheetQueryError("Only a single read-only SELECT statement is allowed.")
    return text


def _clean_sql_error(exc: sqlite3.Error, denied: list[str], tables: list[SheetTable], timeout_seconds: float = QUERY_TIMEOUT_SECONDS) -> str:
    if denied:
        return denied[0]
    message = str(exc).strip() or exc.__class__.__name__
    lowered = message.lower()
    if "one statement at a time" in lowered or "multiple statements" in lowered:
        return "Only a single SELECT statement is allowed (no semicolon-separated statements)."
    if "interrupted" in lowered:
        return f"The query took longer than {timeout_seconds:g} seconds and was stopped. Simplify it (avoid cross joins and unfiltered self-joins)."
    if "not authorized" in lowered:
        return "Only read-only SELECT queries over the spreadsheet tables are allowed."
    if "string or blob too big" in lowered or "too long" in lowered:
        return "The query produced a value that is too large."
    cleaned = f"SQL error: {message}"
    if lowered.startswith(("no such column", "no such table", "ambiguous column")):
        cleaned += f". Available tables and columns: {_describe_tables(tables)}"
    return cleaned


def spreadsheet_query(
    root_path: Path,
    index_path: Path,
    document_ids: list[str],
    sql: str,
    max_rows: Any = None,
    timeout_seconds: float = QUERY_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    if not document_ids:
        raise SpreadsheetQueryError("At least one spreadsheet document is required.")
    statement = _prepare_sql(sql)
    try:
        limit = int(max_rows) if max_rows is not None else DEFAULT_MAX_ROWS
    except (TypeError, ValueError):
        limit = DEFAULT_MAX_ROWS
    limit = max(1, min(limit, HARD_MAX_ROWS))

    tables, skipped, _ = _load_tables(root_path, index_path, document_ids)
    if not tables:
        reason = skipped[0]["reason"] if skipped else "The selected files contain no sheets with data."
        raise SpreadsheetQueryError(f"No spreadsheet data could be loaded. {reason}")

    allowed = {item.table.lower(): item for item in tables}
    denied: list[str] = []
    used: dict[str, SheetTable] = {}
    deadline = [time.monotonic() + timeout_seconds]
    connection = _build_database(tables)
    try:
        _lock_down(connection, allowed, denied, used, deadline)
        deadline[0] = time.monotonic() + timeout_seconds
        try:
            cursor = connection.execute(statement)
            fetched = cursor.fetchmany(limit + 1)
            description = cursor.description or []
        except sqlite3.Error as exc:
            interrupted = "interrupted" in str(exc).lower()
            message = _clean_sql_error(exc, [] if interrupted else denied, tables, timeout_seconds)
            raise SpreadsheetQueryError(message) from None
        columns = [item[0] for item in description]
        truncated = len(fetched) > limit
        rows = [[_json_value(value) for value in row] for row in fetched[:limit]]
    finally:
        connection.close()

    return {
        "columns": columns,
        "rows": rows,
        "rowCount": len(rows),
        "truncated": truncated,
        "tablesUsed": [{"document": item.document, "sheet": item.sheet, "table": item.table} for item in used.values()],
        "skipped": skipped,
        "sql": statement,
    }
