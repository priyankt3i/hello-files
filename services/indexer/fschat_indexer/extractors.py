from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import os
import re
import tempfile
import warnings
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Tuple

# Heavy parser libraries (python-docx, openpyxl, PyMuPDF, xlrd, olefile, pywin32) are
# imported on first use: extraction worker processes then start quickly when a sync only
# touches plain-text files. The import statements stay literal so PyInstaller finds them.


def fs_path(path: Path) -> Path:
    """Add the Windows extended-length prefix when a path approaches or exceeds MAX_PATH (260)."""
    raw = str(path)
    if os.name != "nt" or raw.startswith("\\\\?\\") or len(raw) < 248:
        return path
    abspath = os.path.abspath(raw)
    if abspath.startswith("\\\\"):
        return Path("\\\\?\\UNC\\" + abspath[2:])
    return Path("\\\\?\\" + abspath)


def load_fitz():
    try:
        import fitz  # type: ignore
    except ImportError:
        return None
    return fitz

try:
    from PIL import Image, ImageSequence, UnidentifiedImageError
except ImportError:
    Image = None
    ImageSequence = None
    UnidentifiedImageError = Exception

try:
    import numpy as np
except ImportError:
    np = None

# rapidocr/onnxruntime is imported on first OCR use (see get_ocr_engine): it is the slowest
# import here and most extraction worker processes never need it.
_OCR_UNAVAILABLE = False


def load_xlrd():
    try:
        import xlrd  # type: ignore
    except ImportError:
        return None
    return xlrd


def load_olefile():
    try:
        import olefile  # type: ignore
    except ImportError:
        return None
    return olefile


def load_win32com():
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None
    return win32com


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".ini",
    ".toml",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".css",
    ".scss",
    ".sql",
    ".sh",
    ".ps1",
    ".bat",
    ".cmd",
    ".java",
    ".cs",
    ".go",
    ".rs",
    ".cpp",
    ".c",
    ".h",
    ".hpp",
    ".swift",
    ".kt",
    ".rb",
    ".php",
    ".swl",
    ".svg",
}

WORD_EXTENSIONS = {".docx", ".docm"}
RTF_EXTENSIONS = {".rtf"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm"}
LEGACY_SPREADSHEET_EXTENSIONS = {".xls"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
MIN_TEXT_CHARS_FOR_VISUAL_OCR = 160
MAX_EMBEDDED_IMAGES_PER_PDF_PAGE = 3
MAX_EMBEDDED_IMAGES_PER_DOCX = 12
OCR_MAX_IMAGE_EDGE = 1800
OCR_MAX_IMAGE_PIXELS = 3_000_000
HEADER_SCAN_LIMIT = 12
MAX_PROFILE_HEADERS = 12
SPREADSHEET_GROUP_ROW_THRESHOLD = 100
SPREADSHEET_GROUP_TARGET_CHARS = 1200
MAX_ROW_CHUNKS_PER_SHEET = 250

_OCR_ENGINE = None


def extract_document(path: Path) -> dict[str, Any]:
    path = fs_path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_document(path)
    if suffix in WORD_EXTENSIONS:
        return extract_docx_document(path)
    if suffix in RTF_EXTENSIONS:
        return {"text": extract_rtf(path), "parser": "rtf"}
    if suffix == ".doc":
        return extract_doc_document(path)
    if suffix in SPREADSHEET_EXTENSIONS:
        return extract_workbook_document(path)
    if suffix in LEGACY_SPREADSHEET_EXTENSIONS:
        return extract_legacy_workbook_document(path)
    if suffix == ".csv":
        return {"text": extract_csv(path), "parser": "csv"}
    if suffix in IMAGE_EXTENSIONS:
        return extract_image_file_document(path)
    if suffix in TEXT_EXTENSIONS:
        return {"text": extract_plain_text(path), "parser": "text"}
    text, parser = extract_unknown(path)
    return {"text": text, "parser": parser}


def extract_pdf_document(path: Path) -> dict[str, Any]:
    fitz = load_fitz()
    if fitz is None:
        raise ValueError("PDF support requires PyMuPDF.")

    parts: list[str] = []
    visual_assets: list[dict[str, Any]] = []
    seen_image_xrefs: set[int] = set()
    seen_image_digests: set[str] = set()
    with fitz.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if page_text:
                parts.append(f"# PDF page {page_index}\n{page_text}")

            should_ocr_page = should_attempt_visual_ocr(page_text, MIN_TEXT_CHARS_FOR_VISUAL_OCR)
            if should_ocr_page:
                page_ocr = extract_pdf_page_ocr(page)
                anchor = f"# PDF page {page_index} visual text"
                parts.append(f"{anchor}\n{page_ocr or 'No OCR text detected.'}")
                visual_assets.append(
                    {
                        "kind": "pdf-page",
                        "relativePath": path.name,
                        "label": f"PDF page {page_index}",
                        "anchorText": anchor,
                        "pageNumber": page_index,
                    }
                )

            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                if image_index > MAX_EMBEDDED_IMAGES_PER_PDF_PAGE:
                    break
                xref = int(image_info[0])
                if xref in seen_image_xrefs:
                    continue
                seen_image_xrefs.add(xref)
                image_bytes = document.extract_image(xref).get("image")
                if not image_bytes:
                    continue
                digest = hashlib.sha1(image_bytes).hexdigest()
                if digest in seen_image_digests:
                    continue
                seen_image_digests.add(digest)
                source_name = f"page {page_index} image {image_index}"
                embedded_text = extract_image_bytes(
                    image_bytes,
                    source_name,
                    include_ocr=should_ocr_page,
                )
                if embedded_text:
                    parts.append(embedded_text)
                    visual_assets.append(
                        {
                            **visual_asset_metadata_from_bytes(
                                image_bytes,
                                kind="pdf-image",
                                source_name=source_name,
                                relative_path=path.name,
                            ),
                            "pageNumber": page_index,
                            "imageIndex": image_index,
                        }
                    )

    return {"text": "\n\n".join(part for part in parts if part.strip()), "parser": "pdf", "visualAssets": visual_assets}


def extract_pdf_page_ocr(page) -> str:
    fitz = load_fitz()
    if fitz is None:
        return ""
    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    image_bytes = pixmap.tobytes("png")
    text = ocr_image_bytes(image_bytes)
    return text.strip()


def extract_docx(path: Path) -> str:
    return extract_docx_document(path)["text"]


def extract_docx_document(path: Path) -> dict[str, Any]:
    from docx import Document

    parts: list[str] = []
    visual_assets: list[dict[str, Any]] = []

    try:
        document = Document(path)
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        if paragraph_text:
            parts.append(paragraph_text)
        table_text = extract_docx_tables(document)
        if table_text:
            parts.append(table_text)
    except ValueError:
        xml_paragraphs, xml_tables = extract_docx_xml_parts(path)
        if xml_paragraphs:
            parts.append(xml_paragraphs)
        if xml_tables:
            parts.append(xml_tables)

    native_text = "\n\n".join(parts)
    should_ocr_images = should_attempt_visual_ocr(native_text, MIN_TEXT_CHARS_FOR_VISUAL_OCR)
    embedded_images, embedded_assets = extract_docx_embedded_images(path, include_ocr=should_ocr_images)
    if embedded_images:
        parts.append(embedded_images)
        visual_assets.extend(embedded_assets)

    return {"text": "\n\n".join(part for part in parts if part.strip()), "parser": "docx", "visualAssets": visual_assets}


def extract_docx_xml_parts(path: Path) -> tuple[str, str]:
    """Fallback OpenXML reader for .docm and non-standard .docx content types."""
    w_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs: list[str] = []
    tables: list[str] = []
    table_index = 0
    with zipfile.ZipFile(path) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError("Not a valid Word OpenXML archive.")
        xml_bytes = archive.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    body = root.find(f"{w_ns}body")
    container = body if body is not None else root
    for child in container:
        if child.tag == f"{w_ns}p":
            text = "".join(node.text for node in child.iter(f"{w_ns}t") if node.text).strip()
            if text:
                paragraphs.append(text)
        elif child.tag == f"{w_ns}tbl":
            table_index += 1
            rows: list[str] = [f"# Table {table_index}"]
            for row in child.findall(f".//{w_ns}tr"):
                cells: list[str] = []
                for cell in row.findall(f"./{w_ns}tc"):
                    cell_text = " ".join(
                        "".join(node.text for node in p.iter(f"{w_ns}t") if node.text).strip()
                        for p in cell.findall(f".//{w_ns}p")
                    ).strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    rows.append(" | ".join(cells))
            if len(rows) > 1:
                tables.append("\n".join(rows))
    return "\n".join(paragraphs), "\n".join(tables)


def extract_docx_tables(document: Any) -> str:
    rows: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows.append(f"# Table {table_index}")
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)


def extract_docx_embedded_images(path: Path, include_ocr: bool = True) -> tuple[str, list[dict[str, Any]]]:
    parts: list[str] = []
    visual_assets: list[dict[str, Any]] = []
    seen_digests: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            if len(parts) >= MAX_EMBEDDED_IMAGES_PER_DOCX:
                break
            image_bytes = archive.read(name)
            digest = hashlib.sha1(image_bytes).hexdigest()
            if digest in seen_digests:
                continue
            seen_digests.add(digest)
            image_text = extract_image_bytes(image_bytes, name, include_ocr=include_ocr)
            if image_text:
                parts.append(image_text)
                visual_assets.append(
                    {
                        **visual_asset_metadata_from_bytes(
                            image_bytes,
                            kind="docx-image",
                            source_name=name,
                            relative_path=path.name,
                        ),
                        "mediaPath": name,
                    }
                )
    return "\n\n".join(parts), visual_assets


def extract_doc_document(path: Path) -> dict[str, Any]:
    with path.open("rb") as handle:
        header = handle.read(16)
    if header.lstrip().startswith(b"{\\rtf"):
        return {"text": extract_rtf(path), "parser": "rtf"}
    if header.startswith(b"PK\x03\x04"):
        return extract_docx_document(path)
    return {"text": extract_doc(path), "parser": "doc"}


def extract_doc(path: Path) -> str:
    path = fs_path(path)
    with path.open("rb") as handle:
        header = handle.read(16)
    if header.lstrip().startswith(b"{\\rtf"):
        return extract_rtf(path)
    if header.startswith(b"PK\x03\x04"):
        return extract_docx(path)

    fallback_text = extract_doc_ole_text(path)
    if fallback_text:
        return fallback_text

    if os.environ.get("FSCHAT_ENABLE_WORD_COM") == "1":
        converted_text = extract_doc_via_word(path)
        if converted_text:
            return converted_text

    raise ValueError("Legacy .doc extraction requires readable OLE text streams.")


_RTF_SKIP_DESTINATIONS = {
    "fonttbl",
    "colortbl",
    "stylesheet",
    "info",
    "pict",
    "shppict",
    "nonshppict",
    "object",
    "objdata",
    "themedata",
    "datastore",
    "colorschememapping",
    "latentstyles",
    "rsidtbl",
    "generator",
    "xmlnstbl",
    "listtable",
    "listoverridetable",
    "revtbl",
    "pgdsctbl",
    "mmathPr",
    "defchp",
    "defpap",
}


def extract_rtf(path: Path) -> str:
    raw = fs_path(path).read_bytes()
    embedded_pngs = len(re.findall(rb"89504e470d0a1a0a", raw[:8_000_000], re.IGNORECASE))
    embedded_jpegs = len(re.findall(rb"ffd8ffe0", raw[:8_000_000], re.IGNORECASE))
    compact = re.sub(rb"[0-9a-fA-F\r\n]{256,}", b"", raw).decode("latin-1", errors="ignore")
    text = parse_rtf_text(compact)
    if text:
        return text
    total_images = embedded_pngs + embedded_jpegs
    if total_images > 0:
        return f"# Rich Text Document: {path.name}\nContains {total_images} embedded image(s) with no body text."
    return ""


def parse_rtf_text(rtf: str) -> str:
    out: list[str] = []
    stack: list[bool] = []
    skip = False
    i = 0
    n = len(rtf)
    while i < n:
        ch = rtf[i]
        if ch == "{":
            stack.append(skip)
            i += 1
            if i < n and rtf[i : i + 2] == "\\*":
                skip = True
                i += 2
            continue
        if ch == "}":
            if stack:
                skip = stack.pop()
            i += 1
            continue
        if ch == "\\":
            i += 1
            if i >= n:
                break
            nxt = rtf[i]
            if nxt in "\\{}":
                if not skip:
                    out.append(nxt)
                i += 1
                continue
            if nxt in "\r\n":
                if not skip:
                    out.append("\n")
                i += 1
                continue
            if nxt == "'":
                hex_code = rtf[i + 1 : i + 3]
                if not skip and len(hex_code) == 2:
                    try:
                        out.append(bytes.fromhex(hex_code).decode("cp1252", errors="ignore"))
                    except ValueError:
                        pass
                i += 3
                continue
            if nxt in "~_-":
                if not skip and nxt == "~":
                    out.append(" ")
                elif not skip and nxt in "_-":
                    out.append("-")
                i += 1
                continue
            start = i
            while i < n and rtf[i].isalpha():
                i += 1
            word = rtf[start:i]
            arg_start = i
            if i < n and rtf[i] in "-+":
                i += 1
            while i < n and rtf[i].isdigit():
                i += 1
            arg = rtf[arg_start:i]
            if i < n and rtf[i] == " ":
                i += 1
            if word in _RTF_SKIP_DESTINATIONS:
                skip = True
            elif not skip:
                if word in {"par", "line", "sect", "page", "row"}:
                    out.append("\n")
                elif word in {"tab", "cell"}:
                    out.append("\t")
                elif word == "u" and arg:
                    try:
                        code = int(arg)
                        if code < 0:
                            code += 65536
                        out.append(chr(code))
                    except ValueError:
                        pass
                    if i < n and rtf[i] == "?":
                        i += 1
            continue
        if ch not in "\r\n" and not skip:
            out.append(ch)
        i += 1
    lines = [" ".join(line.split()) for line in "".join(out).splitlines()]
    return "\n".join(line for line in lines if line)


def extract_doc_via_word(path: Path) -> str:
    win32com = load_win32com()
    if win32com is None:
        return ""

    word = None
    document = None
    temp_docx = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(path), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
            temp_docx = Path(handle.name)
        document.SaveAs(str(temp_docx), FileFormat=16)
        return extract_docx(temp_docx)
    except Exception:
        return ""
    finally:
        try:
            if document is not None:
                document.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        if temp_docx and temp_docx.exists():
            temp_docx.unlink(missing_ok=True)


_OLE_BINARY_STREAMS = {"data", "objectpool", "\x01compobj", "\x05summaryinformation", "\x05documentsummaryinformation"}
_OLE_STREAM_READ_LIMIT = 4 * 1024 * 1024


def extract_doc_ole_text(path: Path) -> str:
    olefile = load_olefile()
    if olefile is None:
        return ""

    path_str = str(fs_path(path))
    if not olefile.isOleFile(path_str):
        return ""

    parts: list[str] = []
    with olefile.OleFileIO(path_str) as ole:
        streams = ole.listdir()
        has_word_stream = any(len(s) == 1 and s[0].lower() == "worddocument" for s in streams)
        for stream_name in streams:
            leaf = stream_name[-1].lower() if stream_name else ""
            if leaf in _OLE_BINARY_STREAMS or (has_word_stream and stream_name[0].lower() == "objectpool"):
                continue
            try:
                raw = ole.openstream(stream_name).read(_OLE_STREAM_READ_LIMIT)
            except Exception:
                continue
            extracted = extract_printable_text(raw)
            if extracted:
                parts.append(f"# {'/'.join(stream_name)}\n{extracted}")
    return "\n\n".join(parts)


def extract_legacy_workbook_document(path: Path) -> dict[str, Any]:
    xlrd = load_xlrd()
    if xlrd is None:
        raise ValueError("Legacy .xls support requires xlrd.")

    workbook = xlrd.open_workbook(str(fs_path(path)))
    sheets = []
    for sheet in workbook.sheets():
        rows: list[tuple[int, list[str]]] = []
        for row_index in range(sheet.nrows):
            values = [normalize_spreadsheet_cell(sheet.cell_value(row_index, column_index)) for column_index in range(sheet.ncols)]
            trimmed = trim_trailing_empty_cells(values)
            if any(trimmed):
                rows.append((row_index + 1, trimmed))
        sheets.append({"title": sheet.name, "rows": rows})
    return build_spreadsheet_document(sheets, "spreadsheet-xls")


def extract_workbook_document(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Conditional Formatting extension is not supported and will be removed",
            category=UserWarning,
            module=r"openpyxl\.worksheet\._reader",
        )
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported and will be removed",
            category=UserWarning,
            module=r"openpyxl\.worksheet\._reader",
        )
        warnings.filterwarnings(
            "ignore",
            message="Print area cannot be set",
            category=UserWarning,
            module=r"openpyxl\.reader\.workbook",
        )
        workbook = load_workbook(fs_path(path), read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows: list[tuple[int, list[str]]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [normalize_spreadsheet_cell(value) for value in row]
                trimmed = trim_trailing_empty_cells(values)
                if any(trimmed):
                    rows.append((row_index, trimmed))
            sheets.append({"title": sheet.title, "rows": rows})
    finally:
        # A read-only workbook keeps the file open; on Windows that locks the .xlsx until garbage collection.
        workbook.close()
    return build_spreadsheet_document(sheets, "spreadsheet")


def build_spreadsheet_document(sheets: list[dict[str, Any]], parser_name: str) -> dict[str, Any]:
    parts: list[str] = []
    row_chunks: list[dict[str, Any]] = []
    sheet_profiles: list[dict[str, Any]] = []
    column_hints: list[str] = []

    for sheet in sheets:
        title = str(sheet.get("title") or "Sheet")
        rows: list[tuple[int, list[str]]] = sheet.get("rows") or []
        parts.append(f"# Sheet: {title}")
        if not rows:
            sheet_profiles.append({"name": title, "rowCount": 0, "columnCount": 0, "headerHints": []})
            continue

        header_index = detect_header_row(rows)
        header_values = sanitize_header_values(rows[header_index][1]) if header_index is not None else []
        if header_values:
            parts.append(f"# Header: {' | '.join(header_values)}")
            column_hints.extend(header_values)

        data_row_count = 0
        max_column_count = 0
        sheet_row_entries: list[tuple[int, str]] = []
        for entry_index, (row_number, values) in enumerate(rows):
            max_column_count = max(max_column_count, len(values))
            if header_index is not None and entry_index == header_index:
                continue

            if header_values:
                row_text = build_labeled_row_text(title, row_number, header_values, values)
            else:
                row_text = build_generic_row_text(title, row_number, values)

            sheet_row_entries.append((row_number, row_text))
            if data_row_count < MAX_ROW_CHUNKS_PER_SHEET:
                parts.append(row_text)
            data_row_count += 1

        if len(sheet_row_entries) <= SPREADSHEET_GROUP_ROW_THRESHOLD:
            for row_number, row_text in sheet_row_entries:
                row_chunks.append(
                    {
                        "text": row_text,
                        "chunkType": "spreadsheet-row",
                        "sheetName": title,
                        "rowNumber": row_number,
                    }
                )
        else:
            row_chunks.extend(pack_spreadsheet_row_chunks(title, sheet_row_entries))

        sheet_profiles.append(
            {
                "name": title,
                "rowCount": data_row_count,
                "columnCount": max(max_column_count, len(header_values)),
                "headerHints": header_values[:MAX_PROFILE_HEADERS],
            }
        )

    unique_sheet_names = [profile["name"] for profile in sheet_profiles]
    unique_column_hints = dedupe_preserve_order([item for item in column_hints if item])[:MAX_PROFILE_HEADERS]
    structure = {
        "kind": "spreadsheet",
        "sheetCount": len(sheet_profiles),
        "rowCount": sum(int(profile["rowCount"]) for profile in sheet_profiles),
        "sheetNames": unique_sheet_names,
        "columnHints": unique_column_hints,
        "sheets": sheet_profiles,
    }

    return {
        "text": "\n\n".join(part for part in parts if part.strip()),
        "parser": parser_name,
        "chunks": row_chunks,
        "structure": structure,
    }


def pack_spreadsheet_row_chunks(title: str, row_entries: list[tuple[int, str]]) -> list[dict[str, Any]]:
    """Group consecutive rows on large sheets into ~1.2 KB chunks capped at MAX_ROW_CHUNKS_PER_SHEET."""
    prefix = f"# Sheet: {title}\n"
    total = len(row_entries)
    min_rows_per_chunk = max(1, (total + MAX_ROW_CHUNKS_PER_SHEET - 1) // MAX_ROW_CHUNKS_PER_SHEET)
    chunks: list[dict[str, Any]] = []
    group_texts: list[str] = []
    group_chars = 0
    first_row = 0

    for row_number, row_text in row_entries:
        body = row_text[len(prefix) :] if row_text.startswith(prefix) else row_text
        if group_texts and len(group_texts) >= min_rows_per_chunk and group_chars + len(body) > SPREADSHEET_GROUP_TARGET_CHARS and len(chunks) + 1 < MAX_ROW_CHUNKS_PER_SHEET:
            chunks.append(
                {
                    "text": f"# Sheet: {title}\n\n" + "\n\n".join(group_texts),
                    "chunkType": "spreadsheet-row",
                    "sheetName": title,
                    "rowNumber": first_row,
                }
            )
            group_texts = []
            group_chars = 0
        if not group_texts:
            first_row = row_number
        if group_chars < SPREADSHEET_GROUP_TARGET_CHARS * 2:
            group_texts.append(body)
            group_chars += len(body)

    if group_texts:
        chunks.append(
            {
                "text": f"# Sheet: {title}\n\n" + "\n\n".join(group_texts),
                "chunkType": "spreadsheet-row",
                "sheetName": title,
                "rowNumber": first_row,
            }
        )
    return chunks


def normalize_spreadsheet_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value).strip()


def trim_trailing_empty_cells(values: list[str]) -> list[str]:
    end = len(values)
    while end > 0 and not values[end - 1]:
        end -= 1
    return values[:end]


def dedupe_preserve_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        ordered.append(item)
    return ordered


def detect_header_row(rows: list[tuple[int, list[str]]]) -> int | None:
    best_index: int | None = None
    best_score = 0.0
    for index, (_, values) in enumerate(rows[:HEADER_SCAN_LIMIT]):
        non_empty_values = [value for value in values if value]
        if len(non_empty_values) < 2:
            continue
        score = score_header_candidate(non_empty_values)
        if score > best_score:
            best_score = score
            best_index = index
    return best_index if best_score >= 2.5 else None


def score_header_candidate(values: list[str]) -> float:
    score = 0.0
    unique_values = {value.lower() for value in values if value}
    if len(unique_values) == len(values):
        score += 0.8
    for value in values:
        token_count = len(re.findall(r"[A-Za-z]+", value))
        digit_count = len(re.findall(r"\d", value))
        if token_count > 0:
            score += 1.0
        if digit_count == 0:
            score += 0.35
        if len(value) <= 40:
            score += 0.2
        if re.search(r"[A-Za-z]", value) and not re.fullmatch(r"[A-Z]{2,}\d+", value):
            score += 0.2
    return score


def sanitize_header_values(values: list[str]) -> list[str]:
    headers: list[str] = []
    used: set[str] = set()
    for index, value in enumerate(values, start=1):
        normalized = re.sub(r"\s+", " ", value).strip(" |:-")
        candidate = normalized or f"Column {index}"
        if candidate.lower() in used:
            candidate = f"{candidate} ({index})"
        used.add(candidate.lower())
        headers.append(candidate)
    return headers


def build_labeled_row_text(sheet_name: str, row_number: int, headers: list[str], values: list[str]) -> str:
    parts = [f"# Sheet: {sheet_name}", f"Row: {row_number}"]
    for index, value in enumerate(values, start=1):
        if not value:
            continue
        header = headers[index - 1] if index - 1 < len(headers) else f"Column {index}"
        parts.append(f"{header}: {value}")
    return "\n".join(parts)


def build_generic_row_text(sheet_name: str, row_number: int, values: list[str]) -> str:
    parts = [f"# Sheet: {sheet_name}", f"Row: {row_number}"]
    for index, value in enumerate(values, start=1):
        if not value:
            continue
        parts.append(f"Column {index}: {value}")
    return "\n".join(parts)


def extract_csv(path: Path) -> str:
    rows: list[str] = []
    with fs_path(path).open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            values = [cell.strip() for cell in row if cell and cell.strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)


def extract_plain_text(path: Path) -> str:
    target = fs_path(path)
    try:
        return target.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return target.read_text(encoding="latin-1", errors="ignore")


def extract_image_document(path: Path) -> str:
    if Image is None:
        raise ValueError("Image support requires Pillow.")
    with Image.open(fs_path(path)) as image:
        return describe_image(image, source_name=path.name)


def extract_image_file_document(path: Path) -> dict[str, Any]:
    if Image is None:
        raise ValueError("Image support requires Pillow.")
    image_bytes = fs_path(path).read_bytes()
    with Image.open(io.BytesIO(image_bytes)) as image:
        text = describe_image(image, source_name=path.name)
    return {
        "text": text,
        "parser": "image-ocr",
        "visualAssets": [
            visual_asset_metadata_from_bytes(
                image_bytes,
                kind="image-file",
                source_name=path.name,
                relative_path=path.name,
            )
        ],
    }


def extract_image_bytes(image_bytes: bytes, source_name: str, include_ocr: bool = True) -> str:
    if Image is None:
        return ""
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            return describe_image(image, source_name=source_name, include_ocr=include_ocr)
    except UnidentifiedImageError:
        return ""


def describe_image(image, source_name: str, include_ocr: bool = True) -> str:
    metadata = [f"# Visual asset: {source_name}"]
    metadata.append(f"Format: {getattr(image, 'format', None) or 'unknown'}")
    metadata.append(f"Size: {image.width}x{image.height}")

    frame_count = 1
    if ImageSequence is not None:
        try:
            frame_count = sum(1 for _ in ImageSequence.Iterator(image))
        except Exception:
            frame_count = 1
    if frame_count > 1:
        metadata.append(f"Frames: {frame_count}")

    if include_ocr:
        ocr_text = ocr_image(image)
        if ocr_text:
            metadata.append("")
            metadata.append("# OCR text")
            metadata.append(ocr_text)
        else:
            metadata.append("")
            metadata.append("No OCR text detected.")

    return "\n".join(metadata).strip()


def visual_asset_metadata_from_bytes(image_bytes: bytes, kind: str, source_name: str, relative_path: str) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "kind": kind,
        "relativePath": relative_path,
        "label": source_name,
        "anchorText": f"# Visual asset: {source_name}",
    }
    if Image is None:
        return payload
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            payload["width"] = int(image.width)
            payload["height"] = int(image.height)
            mime_type = Image.MIME.get(getattr(image, "format", "") or "")
            if mime_type:
                payload["mimeType"] = mime_type
    except Exception:
        pass
    return payload


def ocr_image(image) -> str:
    engine = get_ocr_engine()
    if engine is None or np is None:
        return ""

    try:
        prepared = resize_image_for_ocr(image)
        rgb = prepared.convert("RGB")
        result, _ = engine(np.array(rgb))
    except Exception:
        return ""

    if not result:
        return ""

    lines: list[str] = []
    for item in result:
        text = ""
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            candidate = item[1]
            if isinstance(candidate, (list, tuple)) and candidate:
                text = str(candidate[0]).strip()
            else:
                text = str(candidate).strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def ocr_image_bytes(image_bytes: bytes) -> str:
    if Image is None:
        return ""
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            return ocr_image(image)
    except Exception:
        return ""


def get_ocr_engine():
    global _OCR_ENGINE, _OCR_UNAVAILABLE
    if _OCR_ENGINE is None and not _OCR_UNAVAILABLE:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError:
            _OCR_UNAVAILABLE = True
            return None
        try:
            _OCR_ENGINE = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)
        except TypeError:
            _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def extract_unknown(path: Path) -> Tuple[str, str]:
    mime_type, _ = mimetypes.guess_type(path.name)
    if mime_type and mime_type.startswith("text"):
        return extract_plain_text(path), f"mime:{mime_type}"
    if mime_type and mime_type.startswith("image"):
        return extract_image_document(path), f"mime:{mime_type}"
    if path.suffix.lower() in IMAGE_EXTENSIONS:
        return extract_image_document(path), "image-ocr"

    raw = fs_path(path).read_bytes()
    if looks_like_text(raw):
        return raw.decode("utf-8", errors="ignore"), "sniffed-text"

    raise ValueError("Binary or unsupported file type.")


def extract_printable_text(raw: bytes) -> str:
    cleaned: list[str] = []
    seen: set[str] = set()

    for encoding in ("utf-16-le", "latin-1"):
        try:
            decoded = raw.decode(encoding, errors="ignore")
        except Exception:
            continue
        for match in re.finditer(r"[\w][\w\s,\.\-:/\\]{6,}", decoded):
            normalized = " ".join(match.group(0).split())
            if normalized and normalized not in seen:
                seen.add(normalized)
                cleaned.append(normalized)
                if len(cleaned) >= 500:
                    return "\n".join(cleaned)
    return "\n".join(cleaned)


def looks_like_text(raw: bytes) -> bool:
    if not raw:
        return False
    if b"\x00" in raw:
        return False
    sample = raw[:4096]
    printable = sum(1 for byte in sample if 9 <= byte <= 13 or 32 <= byte <= 126)
    ratio = printable / max(len(sample), 1)
    return ratio > 0.75


def normalize_text(text: str) -> str:
    stripped = text.replace("\u0000", "").strip()
    if not stripped:
        return ""
    if stripped[0] in "{[":
        try:
            parsed = json.loads(stripped)
            if isinstance(parsed, (dict, list)):
                return json.dumps(parsed, indent=2, ensure_ascii=True)
        except Exception:
            pass
    return stripped


def should_attempt_visual_ocr(text: str, minimum_chars: int) -> bool:
    normalized = re.sub(r"\s+", "", text or "")
    return len(normalized) < minimum_chars


def resize_image_for_ocr(image):
    if Image is None:
        return image

    width, height = image.size
    longest_edge = max(width, height)
    pixel_count = width * height
    if longest_edge <= OCR_MAX_IMAGE_EDGE and pixel_count <= OCR_MAX_IMAGE_PIXELS:
        return image

    edge_scale = OCR_MAX_IMAGE_EDGE / max(longest_edge, 1)
    pixel_scale = (OCR_MAX_IMAGE_PIXELS / max(pixel_count, 1)) ** 0.5
    scale = min(edge_scale, pixel_scale, 1.0)
    if scale >= 1.0:
        return image

    new_size = (max(1, int(width * scale)), max(1, int(height * scale)))
    resampling_module = getattr(Image, "Resampling", Image)
    return image.resize(new_size, resampling_module.LANCZOS)
